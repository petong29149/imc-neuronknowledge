"""PPTX → one Section per slide with python-pptx (plan T6): "## สไลด์ {n}: {title}", the text of
every shape in (top, left) order — groups included — tables as Markdown, then speaker notes."""

import io

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

from knowledge.converters.base import (
    ConvertResult,
    Section,
    core_properties_metadata,
    markdown_table,
    text_quality,
)
from knowledge.errors import PermanentError
from knowledge.thai_text import DEFAULT_GARBLED_THRESHOLD, normalize


def _shapes(shapes):
    """Every shape, descending into groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _position(shape) -> tuple[int, int]:
    return (shape.top or 0, shape.left or 0)


def _shape_text(shape) -> str:
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        return markdown_table(rows)
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        lines = [p.text.strip() for p in shape.text_frame.paragraphs]
        return "\n".join(line for line in lines if line)
    return ""


def convert_pptx(data: bytes, *, garbled_threshold: float = DEFAULT_GARBLED_THRESHOLD) -> ConvertResult:
    try:
        presentation = pptx.Presentation(io.BytesIO(data))
    except Exception:
        raise PermanentError("เปิดไฟล์ PPTX ไม่ได้ ไฟล์อาจเสียหาย") from None

    sections: list[Section] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = " ".join(title_shape.text_frame.text.split()) if title_shape is not None else ""
        title_id = title_shape.shape_id if title_shape is not None else None

        body = [
            text
            for shape in sorted(_shapes(slide.shapes), key=_position)
            if shape.shape_id != title_id and (text := _shape_text(shape))
        ]
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if not (title or body or notes):
            continue

        parts = [f"## สไลด์ {number}: {title}" if title else f"## สไลด์ {number}", *body]
        if notes:
            parts.append(f"### Notes\n{notes}")
        sections.append(Section(normalize("\n\n".join(parts)), slide=number))

    return ConvertResult(
        sections=sections,
        page_count=len(presentation.slides),
        metadata=core_properties_metadata(presentation.core_properties),
        quality_flag=text_quality(sections, garbled_threshold),
    )
