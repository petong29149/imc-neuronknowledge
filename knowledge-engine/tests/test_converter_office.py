import io
from datetime import datetime

import docx
import pptx
import pytest
from pptx.util import Inches

from knowledge.converters import convert
from knowledge.errors import PermanentError

# --------------------------------------------------------------------------- DOCX


def make_docx() -> bytes:
    document = docx.Document()
    props = document.core_properties
    props.title, props.author, props.keywords = "คู่มือพนักงาน", "ฝ่ายบุคคล", "HR"
    props.created = datetime(2024, 1, 15, 10, 30)
    document.add_paragraph("คำนำของคู่มือฉบับนี้")
    document.add_heading("บทที่ 1 การลา", level=1)
    document.add_paragraph("พนักงานมีสิทธิ์ลาตามประเภทดังนี้")
    document.add_paragraph("ลาพักร้อนปีละสิบวัน", style="List Bullet")
    document.add_paragraph("ลาป่วยได้ตามจริง", style="List Bullet")
    document.add_heading("1.1 ลาพักร้อน", level=2)
    document.add_paragraph("ต้องยื่นล่วงหน้าเจ็ดวัน")
    table = document.add_table(rows=3, cols=3)
    for r, row in enumerate([["ประเภท", "วัน", "หมายเหตุ"], ["ลาพักร้อน", "10", "ต่อปี"], ["ลาป่วย", "30", ""]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    table.cell(2, 1).merge(table.cell(2, 2))
    table.cell(2, 1).text = "30 | ตามจริง"
    document.add_heading("บทที่ 2 สวัสดิการ", level=1)
    document.add_paragraph("ประกันสุขภาพกลุ่ม")
    section = document.sections[0]
    section.header.paragraphs[0].text = "หัวกระดาษที่ไม่ต้องการ"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_docx_sections_follow_headings():
    result = convert("docx", make_docx())
    paths = [(s.heading_path, s.is_table) for s in result.sections]
    assert paths == [
        ([], False),
        (["บทที่ 1 การลา"], False),
        (["บทที่ 1 การลา", "1.1 ลาพักร้อน"], False),
        (["บทที่ 1 การลา", "1.1 ลาพักร้อน"], True),
        (["บทที่ 2 สวัสดิการ"], False),
    ]
    assert result.sections[1].text == ("# บทที่ 1 การลา\n\nพนักงานมีสิทธิ์ลาตามประเภทดังนี้\n\n- ลาพักร้อนปีละสิบวัน\n- ลาป่วยได้ตามจริง")
    assert result.sections[2].text.startswith("## 1.1 ลาพักร้อน")
    assert result.page_count is None
    assert result.quality_flag == "ok"


def test_docx_table_is_markdown_with_merged_cells_once():
    table = convert("docx", make_docx()).sections[3].text
    assert table.splitlines() == [
        "| ประเภท | วัน | หมายเหตุ |",
        "|---|---|---|",
        "| ลาพักร้อน | 10 | ต่อปี |",
        "| ลาป่วย | 30 \\| ตามจริง | |",
    ]


def test_docx_core_properties_and_no_page_header():
    result = convert("docx", make_docx())
    assert result.metadata["title"] == "คู่มือพนักงาน"
    assert result.metadata["author"] == "ฝ่ายบุคคล"
    assert result.metadata["keywords"] == "HR"
    assert result.metadata["created"].startswith("2024-01-15T10:30:00")
    assert all("หัวกระดาษ" not in s.text for s in result.sections)


def test_corrupt_docx():
    with pytest.raises(PermanentError, match="DOCX"):
        convert("docx", b"PK\x03\x04 broken")


# --------------------------------------------------------------------------- PPTX


def make_pptx() -> bytes:
    presentation = pptx.Presentation()
    presentation.core_properties.title = "สรุปนโยบาย"
    presentation.core_properties.author = "ทีม HR"

    slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
    slide.shapes.title.text = "สิทธิการลา"
    low = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
    low.text_frame.text = "ข้อความด้านล่าง"
    high = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    high.text_frame.text = "ลาพักร้อนปีละสิบวัน"
    group = slide.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    inner.text_frame.text = "ข้อความในกลุ่ม"
    slide.notes_slide.notes_text_frame.text = "ผู้บรรยายเน้นเรื่องการยื่นล่วงหน้า"

    slide2 = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank, no title
    table = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(1)).table
    for r, row in enumerate([["ประเภท", "วัน"], ["ลาป่วย", "30"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    presentation.slides.add_slide(presentation.slide_layouts[6])  # empty slide
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def test_pptx_one_section_per_slide_in_reading_order():
    result = convert("pptx", make_pptx())
    assert result.page_count == 3
    assert [s.slide for s in result.sections] == [1, 2]  # the empty slide is skipped
    assert result.sections[0].text == (
        "## สไลด์ 1: สิทธิการลา\n\nลาพักร้อนปีละสิบวัน\n\nข้อความในกลุ่ม\n\nข้อความด้านล่าง\n\n### Notes\nผู้บรรยายเน้นเรื่องการยื่นล่วงหน้า"
    )
    assert result.sections[1].text == "## สไลด์ 2\n\n| ประเภท | วัน |\n|---|---|\n| ลาป่วย | 30 |"
    assert result.metadata == {"title": "สรุปนโยบาย", "author": "ทีม HR"} | {
        k: v for k, v in result.metadata.items() if k in ("created", "modified")
    }


def test_corrupt_pptx():
    with pytest.raises(PermanentError, match="PPTX"):
        convert("pptx", b"PK\x03\x04 broken")


def test_unknown_type():
    with pytest.raises(PermanentError):
        convert("xlsx", b"data")
