"""Small valid files of every supported type, generated in code (no binary fixtures in git)."""

import io
import zipfile

import docx
import pptx
import pymupdf

THAI = "นโยบายการลาพักร้อนของพนักงาน"


def make_pdf(text: str = "Leave policy") -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), text)
    return doc.tobytes()


def make_docx(text: str = THAI) -> bytes:
    document = docx.Document()
    document.add_heading("หัวข้อ", level=1)
    document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_pptx(text: str = THAI) -> bytes:
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = text
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def make_zip(names: list[str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name in names:
            zf.writestr(name, "<x/>")
    return buf.getvalue()


# A Windows executable header: what a renamed .exe looks like
EXE = b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff\x00\x00" + bytes(200)


def _picture() -> "pymupdf.Pixmap":
    import pymupdf

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 300, 200), False)
    pix.clear_with(220)  # a light grey "photo of a page"
    return pix


def make_png() -> bytes:
    return _picture().tobytes("png")


def make_jpg() -> bytes:
    return _picture().tobytes("jpeg")
